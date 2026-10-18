"""Clinical-style .docx report with python-docx.

    from thyrocheck.reports.report_generator import generate_report
    path = generate_report(result, patient={"id": "P-001", "name": "...", "age": 45, "sex": "F"},
                           out_dir="outputs/reports")

CLI:
    python -m thyrocheck.reports.report_generator path/to/scan.png --patient-id P-001
"""
from __future__ import annotations

import io
import uuid
from datetime import datetime
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

RISK_COLOR = {"malignant": RGBColor(0xC0, 0x1C, 0x28), "benign": RGBColor(0x1A, 0x7F, 0x37),
              "normal": RGBColor(0x1F, 0x4E, 0x79)}


def _img_stream(arr: np.ndarray) -> io.BytesIO:
    buf = io.BytesIO()
    Image.fromarray(arr).save(buf, format="PNG")
    buf.seek(0)
    return buf


def generate_report(result: dict, patient: dict | None = None, out_dir: str | Path = "outputs/reports",
                    institution: str = "ThyroCheck AI - Research Prototype",
                    disclaimer: str = "Research prototype - not a medical diagnosis.",
                    model_name: str = "FibonacciNet") -> Path:
    patient = patient or {}
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    report_id = f"TC-{datetime.now():%Y%m%d}-{uuid.uuid4().hex[:6].upper()}"

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = doc.add_heading("Thyroid Ultrasound - AI Analysis Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(institution)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # --- patient / study details
    doc.add_heading("Study details", level=1)
    rows = [
        ("Report ID", report_id),
        ("Date", f"{datetime.now():%d %b %Y, %H:%M}"),
        ("Patient ID", str(patient.get("id", "-"))),
        ("Patient name", str(patient.get("name", "-"))),
        ("Age / Sex", f"{patient.get('age', '-')} / {patient.get('sex', '-')}"),
        ("Referring clinician", str(patient.get("clinician", "-"))),
        ("Model", model_name),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = k, v
        cells[0].paragraphs[0].runs[0].bold = True

    # --- result
    doc.add_heading("AI finding", level=1)
    p = doc.add_paragraph()
    p.add_run("Predicted class: ").bold = True
    run = p.add_run(result["label"].upper())
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RISK_COLOR.get(result["label"], RGBColor(0, 0, 0))
    p.add_run(f"   (confidence {result['confidence']:.1%})")
    if result.get("interpretation"):
        doc.add_paragraph(result["interpretation"])

    prob_table = doc.add_table(rows=1, cols=2)
    prob_table.style = "Light List Accent 1"
    prob_table.rows[0].cells[0].text, prob_table.rows[0].cells[1].text = "Class", "Probability"
    for name, prob in sorted(result["probabilities"].items(), key=lambda kv: -kv[1]):
        c = prob_table.add_row().cells
        c[0].text, c[1].text = name, f"{prob:.1%}"

    # --- images
    doc.add_heading("Explainability (Grad-CAM)", level=1)
    doc.add_paragraph("Left: input ultrasound. Right: Grad-CAM heatmap - warmer regions contributed "
                      "most to the predicted class.")
    img_table = doc.add_table(rows=1, cols=2)
    img_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    img_table.rows[0].cells[0].paragraphs[0].add_run().add_picture(_img_stream(result["image"]), width=Inches(2.9))
    if result.get("overlay") is not None:
        img_table.rows[0].cells[1].paragraphs[0].add_run().add_picture(_img_stream(result["overlay"]), width=Inches(2.9))

    # --- notes + disclaimer
    if patient.get("notes"):
        doc.add_heading("Clinical notes", level=1)
        doc.add_paragraph(str(patient["notes"]))
    doc.add_heading("Disclaimer", level=1)
    d = doc.add_paragraph(disclaimer)
    d.runs[0].italic = True
    d.runs[0].font.size = Pt(9)

    doc.add_paragraph("\nReviewed by: ______________________      Signature: ______________")

    path = out_dir / f"{report_id}.docx"
    doc.save(path)
    return path


def main() -> None:
    from thyrocheck.config import base_argparser, load_config, resolve
    from thyrocheck.inference.predictor import Predictor

    ap = base_argparser("Generate a .docx report for one image")
    ap.add_argument("image")
    ap.add_argument("--patient-id", default="-")
    ap.add_argument("--patient-name", default="-")
    args = ap.parse_args()
    cfg = load_config(args.config, args.set)
    res = Predictor.from_config(cfg).predict_path(args.image, with_gradcam=True)
    path = generate_report(res, {"id": args.patient_id, "name": args.patient_name},
                           resolve(cfg.report.output_dir), cfg.report.institution, cfg.report.disclaimer)
    print(f"[report] {path}")


if __name__ == "__main__":
    main()
